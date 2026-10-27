"""Excel 历史数据字典种子导入（T-S5）。

遵循 PRD 修订 B② 原则：导入内容一律进候选队列，不直接生效——
- 模块 sheet 的表结构文档 → 按 (表名, 字段名) 匹配已同步字段 →
  写 asset_annotation(status=PENDING, source_type=IMPORT) 候选，人工确认后生效；
  已有人工标注（MANUAL）或已确认（CONFIRMED）的一律不覆盖。
- 「数据字典定义」sheet → dict 表（source_type=IMPORT），枚举字典是参考数据，
  按 code upsert，保留人工后续修改（再次导入会覆盖 items，界面上可再编辑）。

解析层是纯函数（只吃行值列表），openpyxl 仅出现在脚本层，便于单测。
"""

from __future__ import annotations

import re
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import UTC, datetime
from typing import Any, BinaryIO

from sqlalchemy import select
from sqlalchemy.dialects.postgresql import insert as pg_insert
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.knowledge import AnnotationHistory, AssetAnnotation, Dictionary
from app.models.metadata import ColumnMeta, TableMeta

_IDENT_RE = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")

_NAME_HEADER_ALIASES = {"列名", "字段", "字段名", "列名称"}
_DESC_HEADER_ALIASES = {"字段描述", "字段说明", "字段含义", "描述"}
_TYPE_HEADER_ALIASES = {"字段类型", "数据类型", "类型"}
_REMARK_HEADER_ALIASES = {"备注", "说明"}


@dataclass
class ExcelColumnDoc:
    table_name: str
    table_desc: str | None
    column_name: str
    column_desc: str | None
    column_type: str | None
    remark: str | None


@dataclass
class ExcelDictDoc:
    code: str
    name: str
    items: list[dict[str, Any]] = field(default_factory=list)


@dataclass
class ParsedWorkbook:
    columns: list[ExcelColumnDoc] = field(default_factory=list)
    dicts: list[ExcelDictDoc] = field(default_factory=list)


@dataclass
class ExcelImportReport:
    column_docs: int = 0
    matched_columns: int = 0
    annotations_created: int = 0
    annotations_skipped_existing: int = 0
    unmatched_docs: int = 0
    dicts_imported: int = 0
    unmatched_tables: list[str] = field(default_factory=list)


def load_sheets_from_fileobj(fh: BinaryIO) -> dict[str, list[list[Any]]]:
    """从文件对象读取 xlsx 全部 sheet 的行值。

    openpyxl 延迟导入：解析与导入逻辑本身不依赖它，便于单测。
    """
    import openpyxl  # type: ignore[import-untyped]

    workbook = openpyxl.load_workbook(fh, read_only=True, data_only=True)
    sheets: dict[str, list[list[Any]]] = {}
    for name in workbook.sheetnames:
        worksheet = workbook[name]
        sheets[name] = [list(row) for row in worksheet.iter_rows(values_only=True)]
    return sheets


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _first_non_empty(texts: list[str], start: int, end: int | None = None) -> str | None:
    stop = end if end is not None else len(texts)
    for t in texts[start:stop]:
        if t and t not in {"表名", "表名描述"}:
            return t
    return None


_TABLE_LABELS = {"表名", "表名称", "表名描述"}


def _last_meaningful(texts: list[str], start: int, end: int) -> str | None:
    """从后往前找表名：跳过空串、标签字面量与纯数字序号。

    真实版式常在表名前带序号或「表名称」标签：
    `1 | 表名称 | tx_order_info | 表名描述 | 订单表`，取第一个非空会抓到序号。
    """
    for t in reversed(texts[start:end]):
        if t and t not in _TABLE_LABELS and not t.isdigit():
            return t
    return None


def _table_block_start(texts: list[str]) -> tuple[str, str | None] | None:
    """识别表块起始行，返回 (表名, 表名描述)。

    两种版式：`表名 | tx_x | · | 表名描述 | 描述`，
    以及无「表名」字面量的 `· | · | tx_x | 表名描述 | · | 描述`。
    """
    if "表名描述" in texts:
        idx = texts.index("表名描述")
        name = _last_meaningful(texts, 0, idx)
        desc = _first_non_empty(texts, idx + 1)
        if name:
            return name, desc
    if "表名" in texts:
        idx = texts.index("表名")
        name = _first_non_empty(texts, idx + 1)
        if name:
            return name, None
    return None


def _header_map(texts: list[str]) -> dict[str, int] | None:
    """按表头文本定位列索引（各 sheet 列序不固定，禁止按位置硬编码）。"""
    mapping: dict[str, int] = {}
    for i, t in enumerate(texts):
        if t in _NAME_HEADER_ALIASES and "name" not in mapping:
            mapping["name"] = i
        elif t in _DESC_HEADER_ALIASES and "desc" not in mapping:
            mapping["desc"] = i
        elif t in _TYPE_HEADER_ALIASES and "type" not in mapping:
            mapping["type"] = i
        elif t in _REMARK_HEADER_ALIASES and "remark" not in mapping:
            mapping["remark"] = i
    if "name" in mapping and "type" in mapping:
        return mapping
    return None


def _at(texts: list[str], mapping: dict[str, int], key: str) -> str | None:
    idx = mapping.get(key)
    if idx is None or idx >= len(texts):
        return None
    return texts[idx] or None


def parse_table_sheet(rows: list[list[Any]]) -> list[ExcelColumnDoc]:
    docs: list[ExcelColumnDoc] = []
    table_name: str | None = None
    table_desc: str | None = None
    header: dict[str, int] | None = None

    for row in rows:
        texts = [_text(c) for c in row]
        if not any(texts):
            continue

        block = _table_block_start(texts)
        if block is not None:
            table_name, table_desc = block
            header = None
            continue

        hmap = _header_map(texts)
        if hmap is not None:
            header = hmap
            continue

        if table_name and header:
            column_name = _at(texts, header, "name")
            # 仅接受合法标识符，防止把说明文字行误当字段
            if column_name and _IDENT_RE.match(column_name):
                docs.append(
                    ExcelColumnDoc(
                        table_name=table_name,
                        table_desc=table_desc,
                        column_name=column_name,
                        column_desc=_at(texts, header, "desc"),
                        column_type=_at(texts, header, "type"),
                        remark=_at(texts, header, "remark"),
                    )
                )
    return docs


def _looks_like_dict_sheet(rows: list[list[Any]]) -> bool:
    for row in rows[:5]:
        texts = {_text(c) for c in row}
        if "type_code" in texts and "dict_code" in texts:
            return True
    return False


def _header_bands(rows: list[list[Any]]) -> list[tuple[int, int]]:
    """检测左右并排多表块 sheet（如 crm）的列带范围。

    同一 sheet 内表头出现在多个不同「列名」列位 → 每个列位一个带，
    带范围取该列位所有表头映射的最小/最大列号。
    """
    bands: dict[int, tuple[int, int]] = {}
    for row in rows:
        texts = [_text(c) for c in row]
        mapping = _header_map(texts)
        if mapping is None:
            continue
        lo, hi = min(mapping.values()), max(mapping.values())
        current = bands.get(mapping["name"])
        if current is None:
            bands[mapping["name"]] = (lo, hi)
        else:
            bands[mapping["name"]] = (min(lo, current[0]), max(hi, current[1]))
    if len(bands) <= 1:
        return []
    return [bands[name_col] for name_col in sorted(bands)]


def parse_table_sheet_banded(rows: list[list[Any]]) -> list[ExcelColumnDoc]:
    bands = _header_bands(rows)
    if not bands:
        return parse_table_sheet(rows)
    docs: list[ExcelColumnDoc] = []
    seen: set[tuple[str, str]] = set()
    for lo, hi in bands:
        sliced = [row[lo : hi + 1] for row in rows]
        for doc in parse_table_sheet(sliced):
            key = (doc.table_name, doc.column_name)
            if key not in seen:
                seen.add(key)
                docs.append(doc)
    return docs


def parse_dict_sheet(rows: list[list[Any]]) -> list[ExcelDictDoc]:
    header: dict[str, int] | None = None
    grouped: dict[str, ExcelDictDoc] = {}

    for row in rows:
        texts = [_text(c) for c in row]
        if header is None:
            idx = {t: i for i, t in enumerate(texts) if t}
            if "type_code" in idx and "dict_code" in idx:
                header = idx
            continue

        type_code = texts[header["type_code"]] if header["type_code"] < len(texts) else ""
        dict_code = texts[header["dict_code"]] if header["dict_code"] < len(texts) else ""
        if not type_code or not dict_code:
            continue

        doc = grouped.get(type_code)
        if doc is None:
            type_name = (
                texts[header["type_name"]]
                if "type_name" in header and header["type_name"] < len(texts)
                else ""
            )
            doc = ExcelDictDoc(code=type_code, name=type_name or type_code)
            grouped[type_code] = doc

        item: dict[str, Any] = {"value": dict_code}
        if (
            "dict_name" in header
            and header["dict_name"] < len(texts)
            and texts[header["dict_name"]]
        ):
            item["label"] = texts[header["dict_name"]]
        if (
            "parent_dict_code" in header
            and header["parent_dict_code"] < len(texts)
            and texts[header["parent_dict_code"]]
        ):
            item["parent"] = texts[header["parent_dict_code"]]
        doc.items.append(item)

    return list(grouped.values())


def parse_workbook(sheets: dict[str, list[list[Any]]]) -> ParsedWorkbook:
    parsed = ParsedWorkbook()
    for rows in sheets.values():
        if _looks_like_dict_sheet(rows):
            parsed.dicts.extend(parse_dict_sheet(rows))
        else:
            parsed.columns.extend(parse_table_sheet_banded(rows))
    return parsed


def _has_cjk(text_value: str) -> bool:
    return any("\u4e00" <= ch <= "\u9fff" for ch in text_value)


def _doc_score(doc: ExcelColumnDoc) -> tuple[int, int, int]:
    """文档质量分：有类型 > 含中文释义 > 释义更长。"""
    desc = doc.column_desc or ""
    return (1 if doc.column_type else 0, 1 if _has_cjk(desc) else 0, len(desc))


async def import_parsed_workbook(
    session: AsyncSession,
    parsed: ParsedWorkbook,
    *,
    operator_id: int | None = None,
) -> ExcelImportReport:
    report = ExcelImportReport(column_docs=len(parsed.columns))

    # (表名, 字段名) 小写索引 → 已同步字段 URN（同名表跨库时全部匹配）
    col_rows = await session.execute(
        select(ColumnMeta.urn, ColumnMeta.column_name, TableMeta.table_name)
        .select_from(ColumnMeta)
        .join(TableMeta, TableMeta.urn == ColumnMeta.table_urn)
        .where(ColumnMeta.is_deleted.is_(False))
    )
    urn_index: dict[tuple[str, str], list[str]] = defaultdict(list)
    for urn, column_name, table_name in col_rows:
        urn_index[(table_name.lower(), column_name.lower())].append(urn)

    # 同一字段可能出现在多个 sheet/区块（如索引清单被误解析成字段行），
    # 取最优文档：有类型的 > 无类型；释义含中文的 > 纯英文；更长的 > 更短。
    best_doc: dict[tuple[str, str], ExcelColumnDoc] = {}
    for doc in parsed.columns:
        key = (doc.table_name.lower(), doc.column_name.lower())
        current = best_doc.get(key)
        if current is None or _doc_score(doc) > _doc_score(current):
            best_doc[key] = doc

    existing_rows = await session.execute(
        select(
            AssetAnnotation.urn,
            AssetAnnotation.business_meaning,
            AssetAnnotation.source_type,
            AssetAnnotation.status,
        )
    )
    existing = {
        urn: (meaning, source_type, status) for urn, meaning, source_type, status in existing_rows
    }

    now = datetime.now(UTC)
    seen_tables: set[str] = set()

    for key, doc in best_doc.items():
        urns = urn_index.get(key)
        if not urns:
            report.unmatched_docs += 1
            seen_tables.add(doc.table_name)
            continue
        report.matched_columns += len(urns)
        if not doc.column_desc:
            continue
        for urn in urns:
            prev = existing.get(urn)
            if prev is not None and (prev[1] == "MANUAL" or prev[2] == "CONFIRMED"):
                report.annotations_skipped_existing += 1
                continue
            before = (
                {"business_meaning": prev[0], "source_type": prev[1], "status": prev[2]}
                if prev
                else None
            )
            stmt = (
                pg_insert(AssetAnnotation)
                .values(
                    urn=urn,
                    asset_type="COLUMN",
                    business_meaning=doc.column_desc,
                    source_desc=doc.remark,
                    status="PENDING",
                    source_type="IMPORT",
                    created_by=operator_id,
                    updated_by=operator_id,
                    created_at=now,
                    updated_at=now,
                )
                .on_conflict_do_update(
                    index_elements=[AssetAnnotation.urn],
                    set_={
                        "business_meaning": doc.column_desc,
                        "source_desc": doc.remark,
                        "status": "PENDING",
                        "source_type": "IMPORT",
                        "updated_by": operator_id,
                        "updated_at": now,
                    },
                )
            )
            await session.execute(stmt)
            await session.execute(
                pg_insert(AnnotationHistory).values(
                    urn=urn,
                    before_data=before,
                    after_data={
                        "business_meaning": doc.column_desc,
                        "source_desc": doc.remark,
                        "status": "PENDING",
                        "source_type": "IMPORT",
                    },
                    operator_id=operator_id,
                    created_at=now,
                )
            )
            existing[urn] = (doc.column_desc, "IMPORT", "PENDING")
            report.annotations_created += 1

    for dict_doc in parsed.dicts:
        stmt = (
            pg_insert(Dictionary)
            .values(
                code=dict_doc.code,
                name=dict_doc.name,
                items=dict_doc.items,
                source_type="IMPORT",
                created_at=now,
                updated_at=now,
            )
            .on_conflict_do_update(
                index_elements=[Dictionary.code],
                set_={
                    "name": dict_doc.name,
                    "items": dict_doc.items,
                    "source_type": "IMPORT",
                    "updated_at": now,
                },
            )
        )
        await session.execute(stmt)
        report.dicts_imported += 1

    report.unmatched_tables = sorted(seen_tables)[:50]
    await session.flush()
    return report
