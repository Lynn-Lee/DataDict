#!/usr/bin/env python
"""导入历史数据字典 Excel 种子（T-S5）。

解析逻辑在 app/services/excel_import.py（纯函数，可单测）；本脚本只负责
用 openpyxl 读文件 + 连库执行导入。依赖：uv pip install -e ".[excel]"

    python scripts/import_excel_seed.py --file /path/to/数据字典.xlsx [--dry-run]

导入遵循候选队列原则：字段描述写为 PENDING/IMPORT 标注（人工确认后生效，
不覆盖已有 MANUAL/CONFIRMED 标注）；枚举字典按 code upsert（source_type=IMPORT）。
"""

from __future__ import annotations

import argparse
import asyncio
from typing import Any

from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.services.excel_import import import_parsed_workbook, parse_workbook

DEFAULT_DB_URL = "postgresql+asyncpg://postgres:postgres@localhost:15432/datadict"


def load_sheets(path: str) -> dict[str, list[list[Any]]]:
    import openpyxl  # type: ignore[import-untyped]  # 延迟导入：仅脚本层依赖 excel extra

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets: dict[str, list[list[Any]]] = {}
    for name in workbook.sheetnames:
        worksheet = workbook[name]
        sheets[name] = [list(row) for row in worksheet.iter_rows(values_only=True)]
    return sheets


async def run(db_url: str, path: str, *, dry_run: bool) -> None:
    sheets = load_sheets(path)
    parsed = parse_workbook(sheets)
    print(f"sheets: {len(sheets)}")
    print(f"parsed: column_docs={len(parsed.columns)} dicts={len(parsed.dicts)}")
    for dict_doc in parsed.dicts[:5]:
        print(f"  dict: {dict_doc.code} {dict_doc.name} items={len(dict_doc.items)}")
    if dry_run:
        tables = {doc.table_name for doc in parsed.columns}
        print(f"distinct tables in docs: {len(tables)}")
        print("dry-run: 未写库")
        return

    engine = create_async_engine(db_url)
    async with async_sessionmaker(engine, expire_on_commit=False)() as session:
        async with session.begin():
            report = await import_parsed_workbook(session, parsed)
    await engine.dispose()

    print("import report:")
    print(f"  column_docs:                {report.column_docs}")
    print(f"  matched_columns:            {report.matched_columns}")
    print(f"  annotations_created:        {report.annotations_created} (PENDING/IMPORT)")
    print(
        f"  annotations_skipped:        {report.annotations_skipped_existing} (已有人工/已确认标注)"
    )
    print(f"  unmatched_docs:             {report.unmatched_docs}")
    print(f"  dicts_imported:             {report.dicts_imported}")
    if report.unmatched_tables:
        preview = ", ".join(report.unmatched_tables[:10])
        print(f"  unmatched tables (前10):    {preview}")
        print("  提示：未匹配多因该表尚未同步进平台，接入对应数据源后重跑即可增量导入")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--file", required=True, help="数据字典 Excel 路径")
    parser.add_argument("--db-url", default=DEFAULT_DB_URL)
    parser.add_argument("--dry-run", action="store_true", help="只解析不写库")
    args = parser.parse_args()
    asyncio.run(run(args.db_url, args.file, dry_run=args.dry_run))


if __name__ == "__main__":
    main()
